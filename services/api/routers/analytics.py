from fastapi import APIRouter, Depends
from sqlalchemy.orm import Session
from sqlalchemy import func, text
from datetime import datetime, timedelta

from database.session import get_db
from database.models import Product

router = APIRouter()


@router.get("/profitability")
def get_profitability(db: Session = Depends(get_db)):
    """
    Returns per-SKU profitability analysis ranked by restock priority.

    Priority score = margin_percent × units_sold_last_30d
    High margin + high sales = restock first.
    Low margin + low sales = restock last.
    """

    cutoff = datetime.utcnow() - timedelta(days=30)

    # Pull sales data from analytics schema for last 30 days
    sales_query = text("""
        SELECT
            p.id,
            p.sku,
            p.name,
            p.category,
            p.unit_cost,
            p.selling_price,
            p.reorder_qty,
            COALESCE(SUM(f.units_sold), 0)  AS units_sold_30d,
            COALESCE(SUM(f.revenue), 0)     AS revenue_30d
        FROM public.products p
        LEFT JOIN analytics.fact_daily_sales f
            ON f.product_id = p.id
            AND f.sale_date >= :cutoff
        GROUP BY p.id, p.sku, p.name, p.unit_cost, p.selling_price, p.reorder_qty
        ORDER BY p.sku
    """)

    rows = db.execute(sales_query, {"cutoff": cutoff.date()}).fetchall()

    result = []
    for row in rows:
        unit_cost      = float(row.unit_cost or 0)
        selling_price  = float(row.selling_price or 0)
        units_sold     = int(row.units_sold_30d or 0)
        revenue        = float(row.revenue_30d or 0)
        reorder_qty    = int(row.reorder_qty or 0)

        profit_per_unit = selling_price - unit_cost
        margin_percent  = round((profit_per_unit / selling_price * 100), 1) if selling_price > 0 else 0
        total_profit    = round(profit_per_unit * units_sold, 2)
        restock_cost    = round(unit_cost * reorder_qty, 2)
        priority_score  = round(margin_percent * units_sold, 2)

        result.append({
            "sku":              row.sku,
            "name":             row.name,
            "category":         row.category or "Uncategorized",
            "selling_price":    selling_price,
            "unit_cost":        unit_cost,
            "profit_per_unit":  round(profit_per_unit, 2),
            "margin_percent":   margin_percent,
            "units_sold_30d":   units_sold,
            "revenue_30d":      round(revenue, 2),
            "total_profit_30d": total_profit,
            "reorder_qty":      reorder_qty,
            "restock_cost":     restock_cost,
            "priority_score":   priority_score,
        })

    # Sort by priority score — highest first
    result.sort(key=lambda x: x["priority_score"], reverse=True)

    return {
        "period_days": 30,
        "total_skus": len(result),
        "products": result
    }


@router.get("/sales-velocity")
def get_sales_velocity(days: int = 7, db: Session = Depends(get_db)):
    """
    Top and bottom movers by units sold in the last N days.
    """
    cutoff = datetime.utcnow() - timedelta(days=days)

    query = text("""
        SELECT
            p.sku,
            p.name,
            COALESCE(SUM(f.units_sold), 0) AS units_sold,
            COALESCE(SUM(f.revenue), 0)    AS revenue
        FROM public.products p
        LEFT JOIN analytics.fact_daily_sales f
            ON f.product_id = p.id
            AND f.sale_date >= :cutoff
        GROUP BY p.sku, p.name
        ORDER BY units_sold DESC
    """)

    rows = db.execute(query, {"cutoff": cutoff.date()}).fetchall()

    products = [
        {
            "sku":        row.sku,
            "name":       row.name,
            "units_sold": int(row.units_sold),
            "revenue":    float(row.revenue),
        }
        for row in rows
    ]

    return {
        "period_days": days,
        "top_movers":    products[:5],
        "bottom_movers": list(reversed(products[-5:])),
    }

from fastapi.responses import StreamingResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io


def style_header_row(ws, row_num: int, num_cols: int, bg_color: str, font_color: str = "FFFFFF"):
    """Apply header styling to a row."""
    fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    font = Font(bold=True, color=font_color, size=11)
    align = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="334155")
    border = Border(bottom=Side(style="medium", color="334155"))

    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_num, col=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = align
        cell.border = border


def auto_width(ws):
    """Auto-fit column widths based on content."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


@router.get("/export")
def export_to_excel(db: Session = Depends(get_db)):
    from database.models import Product, Inventory, Channel, ForecastResult
    from datetime import datetime, timedelta, date as date_type

    try:
        wb = openpyxl.Workbook()

        DARK_HEADER  = "0F172A"
        GREEN_HEADER = "166534"
        BLUE_HEADER  = "1E3A5F"
        PURPLE_HDR   = "4C1D95"
        ORANGE_HDR   = "92400E"
        LIGHT_RED    = "FEE2E2"
        LIGHT_ORANGE = "FEF3C7"
        LIGHT_GREEN  = "DCFCE7"
        ZEBRA        = "F8FAFC"

        def hdr(ws, row, headers, bg):
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col, value=h)
                cell.font = Font(bold=True, color="FFFFFF", size=10)
                cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
                cell.alignment = Alignment(horizontal="center")
            ws.row_dimensions[row].height = 22

        def title(ws, text_val, num_cols, bg):
            ws.merge_cells(f"A1:{get_column_letter(num_cols)}1")
            c = ws["A1"]
            c.value = text_val
            c.font = Font(bold=True, size=13, color="FFFFFF")
            c.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
            c.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 28

        def aw(ws):
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    try:
                        if cell.value:
                            max_len = max(max_len, len(str(cell.value)))
                    except Exception:
                        pass
                ws.column_dimensions[col_letter].width = min(max_len + 4, 40)

        # ── Sheet 1: Sales History — Shopify ────────────────────────────────
        cutoff = datetime.utcnow() - timedelta(days=180)
        sales_query = text("""
            SELECT f.sale_date, p.sku, p.name, p.category,
                   c.name as channel, f.units_sold, f.revenue, f.avg_selling_price
            FROM analytics.fact_daily_sales f
            JOIN public.products p ON p.id = f.product_id
            JOIN public.channels c ON c.id = f.channel_id
            WHERE f.sale_date >= :cutoff
            ORDER BY c.name, f.sale_date DESC, p.sku
        """)
        all_sales = db.execute(sales_query, {"cutoff": cutoff.date()}).fetchall()

        shopify_sales = [r for r in all_sales if r.channel == "shopify"]
        amazon_sales  = [r for r in all_sales if r.channel == "amazon"]

        headers_sales = ["Date", "SKU", "Product Name", "Category",
                         "Units Sold", "Revenue (₹)", "Avg Price (₹)"]

        for sheet_name, rows in [("Sales - Shopify", shopify_sales), ("Sales - Amazon", amazon_sales)]:
            ws = wb.create_sheet(sheet_name)
            title(ws, f"{sheet_name} — Last 180 Days", len(headers_sales), DARK_HEADER)
            hdr(ws, 2, headers_sales, GREEN_HEADER)
            for i, row in enumerate(rows, 3):
                z = ZEBRA if i % 2 == 0 else "FFFFFF"
                data = [str(row.sale_date), row.sku, row.name, row.category or "",
                        int(row.units_sold or 0),
                        round(float(row.revenue or 0), 2),
                        round(float(row.avg_selling_price or 0), 2)]
                for col, val in enumerate(data, 1):
                    c = ws.cell(row=i, column=col, value=val)
                    c.fill = PatternFill(start_color=z, end_color=z, fill_type="solid")
                    c.alignment = Alignment(horizontal="left" if col in [3, 4] else "center")
            aw(ws)
            ws.freeze_panes = "A3"

        # ── Sheet 3: Inventory Snapshot ─────────────────────────────────────
        ws3 = wb.create_sheet("Inventory Snapshot")
        h3 = ["SKU", "Product Name", "Category", "Channel",
              "On Hand", "Reserved", "Available", "Reorder Point", "Status"]
        title(ws3, f"Inventory Snapshot — {datetime.utcnow().strftime('%Y-%m-%d %H:%M')} UTC", len(h3), DARK_HEADER)
        hdr(ws3, 2, h3, BLUE_HEADER)

        inv_rows = (
            db.query(Product.sku, Product.name, Product.category,
                     Channel.name.label("channel"),
                     Inventory.quantity_on_hand, Inventory.quantity_reserved,
                     Product.reorder_point)
            .join(Inventory, Inventory.product_id == Product.id)
            .join(Channel, Channel.id == Inventory.channel_id)
            .order_by(Product.sku, Channel.name)
            .all()
        )
        for i, row in enumerate(inv_rows, 3):
            avail = row.quantity_on_hand - row.quantity_reserved
            if row.quantity_on_hand == 0:
                status, rc = "STOCKOUT", LIGHT_RED
            elif row.quantity_on_hand <= row.reorder_point:
                status, rc = "LOW", LIGHT_ORANGE
            else:
                status, rc = "OK", ZEBRA if i % 2 == 0 else "FFFFFF"
            data = [row.sku, row.name, row.category or "", row.channel,
                    row.quantity_on_hand, row.quantity_reserved, avail,
                    row.reorder_point, status]
            for col, val in enumerate(data, 1):
                c = ws3.cell(row=i, column=col, value=val)
                c.fill = PatternFill(start_color=rc, end_color=rc, fill_type="solid")
                c.alignment = Alignment(horizontal="left" if col in [2, 3] else "center")
                if col == 9:
                    color = "DC2626" if status == "STOCKOUT" else "D97706" if status == "LOW" else "16A34A"
                    c.font = Font(bold=True, color=color)
        aw(ws3)
        ws3.freeze_panes = "A3"

        # ── Sheet 4: Profitability ───────────────────────────────────────────
        ws4 = wb.create_sheet("Profitability Report")
        h4 = ["ABC", "SKU", "Product Name", "Category", "Sell Price (₹)",
              "Unit Cost (₹)", "Profit/Unit (₹)", "Margin %",
              "Units Sold (30d)", "Total Profit (₹)", "Restock Cost (₹)", "Priority"]
        title(ws4, "Profitability Report — Last 30 Days", len(h4), DARK_HEADER)
        hdr(ws4, 2, h4, PURPLE_HDR)

        cutoff30 = datetime.utcnow() - timedelta(days=30)
        prof_q = text("""
            SELECT p.sku, p.name, p.category, p.unit_cost, p.selling_price, p.reorder_qty,
                   COALESCE(SUM(f.units_sold), 0) AS units_sold_30d,
                   COALESCE(SUM(f.revenue), 0) AS revenue_30d
            FROM public.products p
            LEFT JOIN analytics.fact_daily_sales f
                ON f.product_id = p.id AND f.sale_date >= :cutoff
            GROUP BY p.sku, p.name, p.category, p.unit_cost, p.selling_price, p.reorder_qty
            ORDER BY (p.selling_price - p.unit_cost) DESC
        """)
        prof_rows = db.execute(prof_q, {"cutoff": cutoff30.date()}).fetchall()
        total = len(prof_rows)

        for i, row in enumerate(prof_rows, 3):
            uc = float(row.unit_cost or 0)
            sp = float(row.selling_price or 0)
            us = int(row.units_sold_30d or 0)
            ppu = sp - uc
            margin = round(ppu / sp * 100, 1) if sp > 0 else 0
            tp = round(ppu * us, 2)
            rc2 = round(uc * float(row.reorder_qty or 0), 2)
            pri = round(margin * us, 2)
            rank = i - 3
            abc = "A" if rank < total * 0.3 else "B" if rank < total * 0.7 else "C"
            abc_fill = LIGHT_RED if abc == "A" else LIGHT_ORANGE if abc == "B" else LIGHT_GREEN
            z = ZEBRA if i % 2 == 0 else "FFFFFF"
            data = [abc, row.sku, row.name, row.category or "", sp, uc,
                    round(ppu, 2), f"{margin}%", us, tp, rc2, pri]
            for col, val in enumerate(data, 1):
                c = ws4.cell(row=i, column=col, value=val)
                c.fill = PatternFill(start_color=abc_fill if col == 1 else z,
                                     end_color=abc_fill if col == 1 else z, fill_type="solid")
                c.alignment = Alignment(horizontal="left" if col in [3, 4] else "center")
                if col == 1:
                    c.font = Font(bold=True)
        aw(ws4)
        ws4.freeze_panes = "A3"

        # ── Sheet 5: Forecast Summary ────────────────────────────────────────
        ws5 = wb.create_sheet("Forecast Summary")
        h5 = ["SKU", "Product Name", "Category", "Current Stock",
              "Predicted Daily Units", "Stockout Date", "Days Until Stockout",
              "Reorder Flag", "Reorder Qty", "MAPE %"]
        title(ws5, "Demand Forecast Summary — Prophet ML", len(h5), DARK_HEADER)
        hdr(ws5, 2, h5, ORANGE_HDR)

        from database.models import ForecastResult
        fc_rows = (
            db.query(ForecastResult, Product.sku, Product.name, Product.category,
                     Product.reorder_qty, Inventory.quantity_on_hand)
            .join(Product, Product.id == ForecastResult.product_id)
            .join(Inventory, Inventory.product_id == ForecastResult.product_id)
            .order_by(ForecastResult.stockout_date.asc().nullslast())
            .all()
        )
        seen = {}
        for row in fc_rows:
            key = str(row.ForecastResult.product_id)
            if key not in seen:
                seen[key] = row

        today = date_type.today()
        for i, row in enumerate(seen.values(), 3):
            sd = row.ForecastResult.stockout_date
            days_left = (sd - today).days if sd else None
            reorder = row.ForecastResult.reorder_flag
            rc3 = LIGHT_RED if reorder else LIGHT_ORANGE if (days_left and days_left <= 14) else (ZEBRA if i % 2 == 0 else "FFFFFF")
            data = [row.sku, row.name, row.category or "",
                    row.quantity_on_hand,
                    round(float(row.ForecastResult.predicted_units or 0), 2),
                    str(sd) if sd else "Safe (>30d)",
                    days_left if days_left is not None else ">30",
                    "YES" if reorder else "no",
                    row.reorder_qty,
                    round(float(row.ForecastResult.model_mape), 1) if row.ForecastResult.model_mape else "N/A"]
            for col, val in enumerate(data, 1):
                c = ws5.cell(row=i, column=col, value=val)
                c.fill = PatternFill(start_color=rc3, end_color=rc3, fill_type="solid")
                c.alignment = Alignment(horizontal="left" if col in [2, 3] else "center")
                if col == 8 and val == "YES":
                    c.font = Font(bold=True, color="DC2626")
        aw(ws5)
        ws5.freeze_panes = "A3"

        # Remove default empty sheet
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        filename = f"inventory-brain-{datetime.utcnow().strftime('%Y%m%d-%H%M')}.xlsx"
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )

    except Exception as e:
        import traceback
        print("EXPORT ERROR:", traceback.format_exc())
        raise